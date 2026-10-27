"""Lectura, validación y confirmación transaccional; nunca ejecuta fórmulas."""
import hashlib, io, re, zipfile
from datetime import datetime, date, time
from zoneinfo import ZoneInfo
from collections import Counter
from openpyxl import load_workbook
from django.conf import settings
from django.db import transaction
from .schema import *
from .models import Importation, Audit
class InvalidExcel(ValueError): pass

def cell_value(cell):
    value = cell.value
    if isinstance(value,(datetime,date,time)): return value.isoformat()
    if value is None: return ''
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        if float(value).is_integer():
            if re.fullmatch(r'0+',cell.number_format): return str(int(value)).zfill(len(cell.number_format))
            return str(int(value))
    return str(value)

def parse_date(value):
    if isinstance(value,datetime):
        if value.tzinfo: value = value.astimezone(ZoneInfo(settings.TIME_ZONE))
        return value.date()
    if isinstance(value,date): return value
    for fmt in ['%Y-%m-%d','%Y-%m-%d %H:%M:%S','%Y-%m-%dT%H:%M:%S','%d/%m/%Y','%d/%m/%Y %H:%M:%S']:
        try: return datetime.strptime(str(value).strip(),fmt).date()
        except ValueError: pass
    raise ValueError('Fecha inválida. Use fecha Excel, AAAA-MM-DD o DD/MM/AAAA.')

def preview(content,filename,sheet=None):
    if not filename.lower().endswith('.xlsx'): raise InvalidExcel('Solo se permiten archivos .xlsx.')
    if len(content)>settings.MAX_EXCEL_BYTES: raise InvalidExcel('El archivo excede el límite de tamaño.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist())>settings.MAX_EXCEL_UNCOMPRESSED: raise InvalidExcel('Contenido descomprimido excesivo.')
            if len(archive.infolist())>10000: raise InvalidExcel('Demasiadas entradas en el archivo.')
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
        cached=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    except InvalidExcel: raise
    except Exception as exc: raise InvalidExcel('No se pudo leer la estructura del Excel.') from exc
    try:
        expected={header(h):h for h in HEADERS}; candidates=[]; diagnostics=[]
        for ws in wb:
            if ws.max_column and ws.max_column>1000: continue
            first=next(ws.iter_rows(max_row=1),())
            names=[header(c.value) for c in first]
            duplicates=[k for k,v in Counter(n for n in names if n).items() if v>1]
            missing=[v for k,v in expected.items() if k not in names]
            diagnostics.append(f'{ws.title}: faltantes: {", ".join(missing) or "ninguna"}; duplicados: {", ".join(duplicates) or "ninguno"}')
            if not missing and not duplicates: candidates.append(ws.title)
        result={'filename':filename,'hash':hashlib.sha256(content).hexdigest(),'sheets':candidates,'sheet':sheet,'rows':[],'errors':[],'warnings':[],'skipped':0,'start':None,'end':None}
        if not candidates: result['errors']=diagnostics or ['Archivo vacío.']; return result
        if not sheet and len(candidates)>1: return result
        sheet=sheet or candidates[0]; result['sheet']=sheet
        if sheet not in candidates: result['errors']=['La hoja seleccionada no tiene encabezados válidos.']; return result
        ws=wb[sheet]; names=[header(c.value) for c in next(ws.iter_rows(max_row=1))]; positions={v:names.index(k) for k,v in expected.items()}
        seen=set()
        for rownum,(cells,cache) in enumerate(zip(ws.iter_rows(min_row=2),cached[sheet].iter_rows(min_row=2)),2):
            if rownum-1>settings.MAX_EXCEL_ROWS: result['errors'].append('Se excedió el límite de filas.'); break
            if all(c.value is None or c.value=='' for c in cells): result['skipped']+=1; continue
            original={}; values={}; raw={}
            for h,pos in positions.items():
                c=cells[pos]; original[h]={'valor':cell_value(c),'formato':c.number_format,'tipo':c.data_type}
                if c.data_type=='f':
                    c=cache[pos]
                    if c.value is None: result['errors'].append(f'Fila {rownum}, {h}: fórmula sin valor utilizable.')
                raw[h]=c.value; value=cell_value(c)
                values[h]=value.strip() if h in IDENTIFIERS or h==HEADERS[32] else norm(value)
            sid=values['ID']
            if not sid: result['errors'].append(f'Fila {rownum}, ID: ausente.')
            if sid in seen: result['errors'].append(f'Fila {rownum}, ID: duplicado {sid}.')
            seen.add(sid)
            try: day=parse_date(raw['Fecha Registro'])
            except ValueError as exc: result['errors'].append(f'Fila {rownum}, Fecha Registro: {exc}'); continue
            values['Fecha Registro']=day.isoformat()
            for h in CONTROLS+DEFECTS:
                allowed=['CUMPLE','NO CUMPLE','N/A'] if h in CONTROLS else ['LIMPIO','DETECTADO','N/A']
                if not values[h]: result['warnings'].append(f'Fila {rownum}, {h}: Sin dato.')
                elif values[h] not in allowed: result['errors'].append(f'Fila {rownum}, {h}: valor desconocido «{values[h]}».')
            if values['Estatus'] not in ['CONFORME','NO CONFORME']:
                result['errors'].append(f'Fila {rownum}, Estatus: desconocido «{values["Estatus"]}».')
            failures=sum(values[h]=='NO CUMPLE' for h in CONTROLS); detections=sum(values[h]=='DETECTADO' for h in DEFECTS)
            result['rows'].append(dict(source_id=sid,source_row=rownum,date=day.isoformat(),values=values,original=original,failures=failures,detections=detections,passes=sum(values[h]=='CUMPLE' for h in CONTROLS),missing=sum(not values[h] for h in CONTROLS),na=sum(values[h]=='N/A' for h in CONTROLS),inconsistent=values['Estatus']=='CONFORME' and bool(failures or detections),**{k:values[v] for k,v in DIMENSIONS.items()}))
        if not result['rows']: result['errors'].append('No hay auditorías válidas en la hoja.')
        else:
            dates=[r['date'] for r in result['rows']]; result['start']=min(dates); result['end']=max(dates)
        return result
    finally: wb.close(); cached.close()

@transaction.atomic
def confirm(result,user):
    if result['errors'] or not result['rows']: raise InvalidExcel('No se puede confirmar una importación con errores o sin registros.')
    existing=Importation.objects.filter(sha256=result['hash']).first()
    if existing: return existing,False
    imp=Importation.objects.create(filename=result['filename'],sha256=result['hash'],sheet=result['sheet'],user=user,count=len(result['rows']),start=result['start'],end=result['end'],warnings=result['warnings'])
    Audit.objects.bulk_create([Audit(importation=imp,**r) for r in result['rows']])
    return imp,True
