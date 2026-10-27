import io, copy, tempfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from django.test import TestCase, override_settings
from django.contrib.auth.models import User, Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from openpyxl import Workbook
from .schema import HEADERS, CONTROLS, DEFECTS
from .importer import preview, confirm, InvalidExcel
from .models import Importation, Audit
from .metrics import metrics, filtered
from .views import csv_safe

def row(**changes):
    r={h:'' for h in HEADERS}; r.update({h:'CUMPLE' for h in CONTROLS}); r.update({h:'LIMPIO' for h in DEFECTS})
    r.update({'ID':'001','Fecha Registro':datetime(2026,10,5,23,59),'Turno':'Día','Proceso':'Doblez','No. Parte':'00042','Líneas Auditadas':'L1','Estatus':'CONFORME','Máquinas':'M1','Nomina Inspector':'007'})
    r.update(changes); return r

def excel(rows=None,headers=None,sheets=1):
    wb=Workbook(); wb.remove(wb.active); headers=headers or HEADERS
    for i in range(sheets):
        ws=wb.create_sheet('Datos '+str(i)); ws.append(headers)
        for r in rows if rows is not None else [row()]: ws.append([r.get(h,'') for h in headers])
    out=io.BytesIO(); wb.save(out); return out.getvalue()

class ImportTests(TestCase):
    def setUp(self): self.user=User.objects.create_user('prueba',password='clave-de-prueba-123')
    def parse(self,rows=None,**kwargs): return preview(excel(rows,**kwargs),'sintetico.xlsx')
    def test_valid(self):
        r=self.parse(); self.assertEqual(r['errors'],[]); self.assertEqual(r['start'],'2026-10-05'); self.assertEqual(r['rows'][0]['source_id'],'001')
        imp,new=confirm(r,self.user); self.assertTrue(new); self.assertEqual(imp.audits.count(),1)
    def test_reordered_normalized(self):
        wb=Workbook(); ws=wb.active
        import unicodedata
        headers=list(reversed(HEADERS)); ws.append(['  '+''.join(c for c in unicodedata.normalize('NFD',h.lower()) if not unicodedata.combining(c))+'  ' for h in headers]); ws.append([row()[h] for h in headers]); out=io.BytesIO(); wb.save(out)
        self.assertEqual(preview(out.getvalue(),'x.xlsx')['errors'],[])
    def test_missing_columns(self): self.assertTrue(self.parse(headers=HEADERS[:-1])['errors'])
    def test_duplicate_headers(self): self.assertTrue(self.parse(headers=HEADERS+[' id '])['errors'])
    def test_extra_columns(self): self.assertFalse(self.parse(headers=HEADERS+['Extra'])['errors'])
    def test_invalid_dates(self): self.assertTrue(self.parse([row(**{'Fecha Registro':'10-11-12'})])['errors'])
    def test_numeric_zero_format(self):
        wb=Workbook(); ws=wb.active; ws.append(HEADERS); ws.append([row()[h] for h in HEADERS]); ws.cell(2,5,42).number_format='000000'; out=io.BytesIO();wb.save(out)
        self.assertEqual(preview(out.getvalue(),'x.xlsx')['rows'][0]['parte'],'000042')
    def test_na_blank_unknown(self):
        r=self.parse([row(**{CONTROLS[0]:'N/A',CONTROLS[1]:'',DEFECTS[0]:''})]); self.assertFalse(r['errors']);self.assertEqual(len(r['warnings']),2); self.assertEqual(r['rows'][0]['passes'],9)
        self.assertTrue(self.parse([row(**{CONTROLS[0]:'OK'})])['errors'])
        self.assertTrue(self.parse([row(**{DEFECTS[0]:'NO CUMPLE'})])['errors'])
        self.assertTrue(self.parse([row(Estatus='OTRO')])['errors'])
    def test_empty(self): self.assertTrue(self.parse([])['errors'])
    def test_duplicate_ids(self): self.assertTrue(self.parse([row(),row()])['errors'])
    def test_repeat_file(self):
        r=self.parse(); a,_=confirm(r,self.user); b,new=confirm(r,self.user); self.assertFalse(new); self.assertEqual(a,b)
    def test_invalid_confirmation_no_partial(self):
        r=self.parse([row(),row(**{'ID':'002','Fecha Registro':'mala'})])
        with self.assertRaises(InvalidExcel): confirm(r,self.user)
        self.assertFalse(Importation.objects.exists()); self.assertFalse(Audit.objects.exists())
    def test_transaction_rollback(self):
        r=self.parse(); r['rows'].append(copy.deepcopy(r['rows'][0]))
        with self.assertRaises(IntegrityError): confirm(r,self.user)
        self.assertFalse(Importation.objects.exists())
    def test_multiple_sheets(self):
        r=self.parse(sheets=2); self.assertEqual(len(r['sheets']),2); self.assertFalse(r['rows'])
        r=preview(excel(sheets=2),'x.xlsx','Datos 1');self.assertEqual(len(r['rows']),1)
    def test_formulas(self): self.assertTrue(self.parse([row(**{'ID':'=1+1'})])['errors'])
    @override_settings(MAX_EXCEL_BYTES=1)
    def test_size_limit(self):
        with self.assertRaises(InvalidExcel): self.parse()
    @override_settings(MAX_EXCEL_ROWS=1)
    def test_row_limit(self): self.assertTrue(self.parse([row(),row(ID='2')])['errors'])
    @override_settings(MAX_EXCEL_UNCOMPRESSED=1)
    def test_zip_limit(self):
        with self.assertRaises(InvalidExcel): self.parse()
    def test_skip_only_empty(self):
        r=self.parse([row(),{},row(ID='002')]);self.assertEqual(r['skipped'],1);self.assertEqual(len(r['rows']),2)
    def test_zero_denominators(self):
        imp,_=confirm(self.parse([row(**{h:'N/A' for h in CONTROLS})]),self.user)
        m=metrics(imp.audits.all());self.assertIsNone(m['summary']['compliance']);self.assertEqual(m['summary']['na'],11)
    def test_combined_filters(self):
        imp,_=confirm(self.parse([row(),row(ID='2',Turno='NOCHE',**{CONTROLS[0]:'NO CUMPLE'})]),self.user)
        q=filtered(imp,{'turno':'NOCHE','linea':'L1','desde':'2026-10-05','hallazgos':'controles'}); self.assertEqual(q.count(),1)
        self.assertEqual(filtered(imp,{'turno':'DÍA','hallazgos':'controles'}).count(),0)
    def test_weekly_and_gaps(self):
        imp,_=confirm(self.parse([row(),row(ID='2',**{'Fecha Registro':datetime(2026,10,7)}),row(ID='3',**{'Fecha Registro':datetime(2026,10,19)})]),self.user)
        daily=metrics(imp.audits.all())['timeline']; self.assertIsNone(daily[1]['pct']); self.assertEqual(daily[1]['total'],0)
        weekly=metrics(imp.audits.all(),True)['timeline'];self.assertEqual(weekly[0]['label'],'2026-10-05');self.assertEqual(weekly[0]['total'],2);self.assertIsNone(weekly[1]['pct'])
    def test_A_B_A_every_metric_and_views(self):
        a,_=confirm(self.parse([row(**{CONTROLS[0]:'NO CUMPLE'})]),self.user)
        b,_=confirm(self.parse([row(),row(ID='2',**{DEFECTS[0]:'DETECTADO'})]),self.user)
        expected_a=metrics(a.audits.all());expected_b=metrics(b.audits.all())
        self.assertEqual(expected_a['summary']['failures'],1);self.assertEqual(expected_b['summary']['failures'],0)
        self.user.is_superuser=True;self.user.save();self.client.force_login(self.user)
        for imp,expected in [(a,expected_a),(b,expected_b),(a,expected_a)]:
            r=self.client.get('/',{'importacion':imp.pk},follow=True);self.assertEqual(r.status_code,200);self.assertEqual(r.context['metrics'],expected)
            response=self.client.get('/datos/');self.assertEqual(response.json(),expected)
            response=self.client.get('/exportar/'); self.assertNotIn('\ufeff',response.content.decode('utf-8-sig')); self.assertEqual(len(list(__import__('csv').reader(io.StringIO(response.content.decode('utf-8-sig'))))),imp.count+1)
        self.assertEqual(self.client.get('/auditoria/'+str(b.audits.first().pk)+'/').status_code,404)
    def test_weighted_lines_and_union(self):
        imp,_=confirm(self.parse([row(**{CONTROLS[0]:'NO CUMPLE',DEFECTS[0]:'DETECTADO'}),row(ID='2',**{h:'N/A' for h in CONTROLS})]),self.user)
        m=metrics(imp.audits.all()); self.assertEqual(m['summary']['union'],1); self.assertEqual(m['lines'][0]['applicable'],11);self.assertEqual(m['lines'][0]['compliance'],90.91)
    def test_access(self):
        for url in ['/','/datos/','/exportar/','/importar/','/auditoria/1/']: self.assertEqual(self.client.get(url).status_code,302)
        self.client.force_login(self.user)
        for url in ['/','/datos/','/importar/']:self.assertEqual(self.client.get(url).status_code,200)
        self.assertEqual(self.client.get('/exportar/').status_code,404)
        self.assertEqual(self.client.get('/auditoria/1/').status_code,404)
    def test_upload_preview_confirm(self):
        self.user.is_superuser=True;self.user.save();self.client.force_login(self.user)
        with tempfile.TemporaryDirectory() as tmp, override_settings(BASE_DIR=Path(tmp)):
            response=self.client.post('/importar/',{'archivo':SimpleUploadedFile('A.xlsx',excel())});self.assertEqual(response.status_code,200); r=response.context['result'];self.assertFalse(Importation.objects.exists())
            response=self.client.post('/importar/',{'confirmar':'1','hash':r['hash'],'hoja':r['sheet']});self.assertEqual(response.status_code,302);self.assertEqual(Importation.objects.count(),1)
    def test_csv_injection(self):
        for value in ['=1+1',' +cmd','@SUM(A1)','-1','\t=1']:self.assertTrue(csv_safe(value).startswith("'"))
        self.assertEqual(csv_safe('00042'),"'00042")
    def test_reference(self):
        path=Path(__file__).resolve().parent.parent/'Reporte_Auditorias_FCD_2026_10 (2).xlsx'
        if not path.exists(): self.skipTest('Excel de referencia no disponible')
        r=preview(path.read_bytes(),path.name);self.assertFalse(r['errors']); imp,_=confirm(r,self.user);s=metrics(imp.audits.all())['summary']
        self.assertEqual((imp.count,str(imp.start),str(imp.end)),(149,'2026-10-05','2026-10-26'))
        self.assertEqual((s['with_failures'],s['failures'],s['with_defects']),(95,126,38))
        self.assertEqual(imp.audits.filter(values__Estatus='CONFORME').count(),149)
        for h,n in [('Validación Mtto',85),('Material Identificado',31)]:self.assertEqual(imp.audits.filter(**{'values__'+h:'NO CUMPLE'}).count(),n)

    def test_presentation_permissions_and_empty(self):
        self.assertEqual(self.client.get('/presentacion/').status_code,302)
        self.client.force_login(self.user)
        response=self.client.get('/presentacion/')
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'Aún no hay importaciones')

    def test_presentation_matches_dashboard_filters(self):
        a,_=confirm(self.parse([row(),row(ID='2',Turno='NOCHE',**{CONTROLS[0]:'NO CUMPLE'})]),self.user)
        self.user.is_superuser=True;self.user.save();self.client.force_login(self.user)
        params={'importacion':a.pk,'turno':'NOCHE','hallazgos':'controles','periodo':'semanal'}
        dashboard=self.client.get('/',params)
        presentation=self.client.get('/presentacion/',params)
        self.assertEqual(presentation.status_code,200)
        self.assertEqual(presentation.context['metrics'],dashboard.context['metrics'])
        self.assertEqual(presentation.context['metrics']['summary']['total'],1)
        self.assertContains(presentation,'data-chart=',count=16)
        self.assertContains(dashboard,'/presentacion/?')
        self.assertNotContains(presentation,'Consulta de auditorías')

    def test_presentation_switch_import_resets_filters(self):
        a,_=confirm(self.parse(),self.user)
        b,_=confirm(self.parse([row(ID='2',Turno='NOCHE')]),self.user)
        self.user.is_superuser=True;self.user.save();self.client.force_login(self.user)
        self.client.get('/',{'importacion':a.pk})
        response=self.client.get('/presentacion/',{'importacion':b.pk,'turno':'DÍA'},follow=True)
        self.assertEqual(response.redirect_chain[0][0],'/presentacion/?importacion='+str(b.pk))
        self.assertEqual(response.context['metrics']['summary']['total'],1)

    def expanded_fixture(self):
        rows=[
            row(ID='1',Proceso='P1',**{CONTROLS[0]:'NO CUMPLE',DEFECTS[0]:'DETECTADO',DEFECTS[1]:'DETECTADO'}),
            row(ID='2',Turno='NOCHE',Proceso='P2',**{'Fecha Registro':datetime(2026,10,7),CONTROLS[0]:'NO CUMPLE'}),
            row(ID='3',Proceso='P2',Estatus='NO CONFORME',**{'Fecha Registro':datetime(2026,10,7),'Líneas Auditadas':'L2',DEFECTS[0]:'DETECTADO'}),
            row(ID='4',Proceso='P1',**{'Fecha Registro':datetime(2026,10,12),**{h:'N/A' for h in CONTROLS}}),
            row(ID='5',Turno='NOCHE',Proceso='P2',**{'Fecha Registro':datetime(2026,10,14),'Líneas Auditadas':'L2',**{h:'' for h in CONTROLS}}),
        ]
        return confirm(self.parse(rows),self.user)[0]

    def test_requirement_states_and_distribution(self):
        imp=self.expanded_fixture();m=metrics(imp.audits.all());first=m['requirement_profile'][0]
        self.assertEqual((first['passes'],first['failures'],first['na'],first['missing']),(1,2,1,1))
        self.assertEqual(first['compliance'],33.33)
        for r in m['requirement_profile']:self.assertEqual(r['passes']+r['failures']+r['na']+r['missing'],5)
        dist={r['number']:r for r in m['failure_distribution']}
        self.assertEqual((dist[0]['count'],dist[1]['count'],dist[0]['with_missing']),(3,2,1))

    def test_status_groups_do_not_duplicate_audits(self):
        imp=self.expanded_fixture();rows=metrics(imp.audits.all())['status_findings'];conforme=next(r for r in rows if r['label']=='CONFORME')
        self.assertEqual((conforme['both'],conforme['controls_only'],conforme['visual_only'],conforme['neither'],conforme['missing']),(1,1,0,2,1))
        for r in rows:self.assertEqual(sum(r[k] for k in ['both','controls_only','visual_only','neither']),r['total'])
        self.assertEqual(sum(r['total'] for r in rows),5)

    def test_line_shift_coverage_and_process_detection_counts(self):
        imp=self.expanded_fixture();m=metrics(imp.audits.all());coverage=m['line_shifts']
        self.assertEqual(sum(sum(r['counts']) for r in coverage['rows']),5)
        line=next(r for r in coverage['rows'] if r['label']=='L1');self.assertEqual(line['total'],3)
        p1=next(r for r in m['visual_processes'] if r['label']=='P1')
        self.assertEqual((p1['detections'],p1['with_defects'],p1['total'],p1['pct']),(2,1,2,50))
        self.assertEqual(sum(p1['counts']),p1['detections'])
        cell=next(r for r in m['process_heat'] if r['label']=='P1')['cells'][0]
        self.assertEqual((cell['bad'],cell['den'],cell['pct']),(1,1,100))

    def test_compliance_evolution_weighting_week_and_missing(self):
        imp=self.expanded_fixture();weekly=metrics(imp.audits.all(),True)['timeline'];self.assertEqual((weekly[0]['passes'],weekly[0]['applicable'],weekly[0]['compliance']),(31,33,93.94))
        self.assertIsNone(weekly[1]['compliance'])
        daily=metrics(imp.audits.all())['timeline'];self.assertIsNone(daily[1]['compliance'])
        missing=next(r for r in daily if r['label']=='2026-10-14');self.assertEqual(missing['total'],1);self.assertIsNone(missing['compliance'])

    def test_additional_charts_scope_and_explanations(self):
        from .chart_catalog import CHARTS
        imp=self.expanded_fixture();only=metrics(filtered(imp,{'proceso':'P1'}))
        self.assertEqual(only['summary']['total'],2);self.assertEqual(sum(r['count'] for r in only['failure_distribution']),2)
        self.assertEqual(sum(r['total'] for r in only['visual_processes']),2)
        empty=metrics(filtered(imp,{'proceso':'NO EXISTE'}));self.assertFalse(empty['status_findings']);self.assertFalse(empty['visual_processes']);self.assertTrue(all(r['compliance'] is None for r in empty['requirement_profile']))
        self.assertEqual(len({c['id'] for c in CHARTS}),16)
        self.assertTrue(all(c[key] for c in CHARTS for key in ['title','meaning','calculation','interpretation']))
        self.user.is_superuser=True;self.user.save();self.client.force_login(self.user)
        for url in ['/','/presentacion/']:
            response=self.client.get(url,{'importacion':imp.pk,'proceso':'P1'})
            self.assertContains(response,'class="chart-help-button"',count=16)
            self.assertContains(response,'id="chart-help"',count=1)
            self.assertEqual(response.context['metrics'],only)

    def test_operator_compliance_preserves_identifiers_and_weights(self):
        imp,_=confirm(self.parse([
            row(ID='1',**{'Nómina Operador':'0007','Nomina Inspector':'9999'}),
            row(ID='2',**{'Nómina Operador':'0007',**{h:'N/A' for h in CONTROLS},CONTROLS[0]:'NO CUMPLE'}),
            row(ID='3',**{'Nómina Operador':'7','Nomina Inspector':'0007',**{h:'N/A' for h in CONTROLS}}),
            row(ID='4',**{'Nómina Operador':''})
        ]),self.user)
        groups={r['operador']:r for r in metrics(imp.audits.all())['operators']}
        self.assertEqual(len(groups),3)
        self.assertEqual((groups['0007']['total'],groups['0007']['passes'],groups['0007']['applicable'],groups['0007']['compliance']),(2,11,12,91.67))
        self.assertEqual(groups['0007']['bad'],1);self.assertEqual(groups['0007']['na'],10)
        self.assertIsNone(groups['7']['compliance']);self.assertEqual(groups['']['total'],1)
        self.assertEqual(filtered(imp,{'operador':'0007'}).count(),2)
        self.assertEqual(filtered(imp,{'operador':'0007','inspector':'9999'}).count(),1)
        self.user.is_superuser=True;self.user.save();self.client.force_login(self.user)
        response=self.client.get('/',{'importacion':imp.pk,'operador':'0007'})
        self.assertEqual(response.context['metrics']['summary']['total'],2)
        self.assertContains(response,'Nómina Operador')
        self.assertContains(response,'data-help="operators"')
        self.assertEqual(response.context['page'].paginator.count,2)

    def test_operator_backfill_from_existing_json(self):
        from importlib import import_module
        from django.apps import apps
        from django.db import connection
        imp,_=confirm(self.parse([row(**{'Nómina Operador':'0000042'})]),self.user)
        audit=imp.audits.get();Audit.objects.filter(pk=audit.pk).update(operador='')
        from types import SimpleNamespace
        import_module('audits.migrations.0003_backfill_operator').backfill_operator(apps,SimpleNamespace(connection=connection))
        audit.refresh_from_db();self.assertEqual(audit.operador,'0000042')

    def test_asset_versions_change_and_chart_guides_are_in_headers(self):
        from audits.templatetags.asset_tags import asset
        with tempfile.TemporaryDirectory() as tmp, override_settings(BASE_DIR=Path(tmp)):
            static=Path(tmp)/'static';static.mkdir();file=static/'guide.css'
            file.write_text('a{}');first=asset('guide.css')
            file.write_text('a{color:blue}');self.assertNotEqual(asset('guide.css'),first)
        imp,_=confirm(self.parse(),self.user)
        self.user.is_superuser=True;self.user.save();self.client.force_login(self.user)
        response=self.client.get('/',{'importacion':imp.pk})
        self.assertContains(response,'class="chart-card-heading"',count=16)
        self.assertContains(response,'class="help-control-label"',count=16)
        self.assertContains(response,'/static/chart-help.css?v=')

    def test_configurable_private_upload_directory(self):
        self.user.is_superuser=True;self.user.save();self.client.force_login(self.user)
        with tempfile.TemporaryDirectory() as tmp, override_settings(PRIVATE_DIR=Path(tmp)/'uploads'):
            response=self.client.post('/importar/',{'archivo':SimpleUploadedFile('docker.xlsx',excel())})
            self.assertEqual(response.status_code,200)
            self.assertEqual(len(list((Path(tmp)/'uploads').glob('*.xlsx'))),1)
            result=response.context['result']
            response=self.client.post('/importar/',{'confirmar':'1','hash':result['hash'],'hoja':result['sheet']})
            self.assertEqual(response.status_code,302)
            self.assertFalse(list((Path(tmp)/'uploads').glob('*.xlsx')))

    def test_production_requires_secret(self):
        import subprocess,os,sys
        env=os.environ.copy();env['DEBUG']='0';env['SECRET_KEY']=''
        result=subprocess.run([sys.executable,'-c','import config.settings'],env=env,cwd=Path(__file__).resolve().parent.parent,capture_output=True,text=True)
        self.assertNotEqual(result.returncode,0)
        self.assertIn('Configure SECRET_KEY',result.stderr)

    def test_production_static_collection(self):
        from django.core.management import call_command
        with tempfile.TemporaryDirectory() as tmp, override_settings(DEBUG=False,STATIC_ROOT=Path(tmp)):
            call_command('collectstatic',interactive=False,verbosity=0)
            self.assertTrue((Path(tmp)/'chart-help.css').is_file())
            self.assertTrue((Path(tmp)/'bar-track.js').is_file())
